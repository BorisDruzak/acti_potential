import csv
import io
from uuid import uuid4

from tests.test_documents import headers, setup_registry


def cards(client, factory):
    registry = setup_registry(client, factory)
    result = []
    for title in ['Selected first', 'Other second', 'Selected third']:
        response = client.post('/api/v1/documents', json={
            'registry_id': registry['id'], 'document_type_id': registry['document_type_ids'][0],
            'values': {'title': title},
        }, headers=headers(client))
        assert response.status_code == 201, response.text
        result.append(response.json())
    return result


def test_selection_export_is_exact_and_keeps_requested_order(client, factory):
    first, second, third = cards(client, factory)
    response = client.post('/api/v1/documents/export-selection', json={
        'document_ids': [third['id'], first['id']], 'format': 'csv',
        'columns': ['title'], 'mode': 'editor',
    }, headers=headers(client))
    assert response.status_code == 200, response.text
    rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert rows == [['Наименование документа [title]'], ['Selected third'], ['Selected first']]
    assert second['title'] not in response.text


def test_selection_export_never_silently_skips_an_unavailable_id(client, factory):
    first, _, _ = cards(client, factory)
    response = client.post('/api/v1/documents/export-selection', json={
        'document_ids': [first['id'], str(uuid4())], 'format': 'xlsx', 'mode': 'editor',
    }, headers=headers(client))
    assert response.status_code == 404, response.text
    assert response.headers['content-type'].startswith('application/json')


def test_selection_export_rejects_duplicates_empty_and_missing_csrf(client, factory):
    first, _, _ = cards(client, factory)
    for ids in [[], [first['id'], first['id']]]:
        response = client.post('/api/v1/documents/export-selection', json={
            'document_ids': ids, 'format': 'csv',
        }, headers=headers(client))
        assert response.status_code == 422, response.text
    assert client.post('/api/v1/documents/export-selection', json={
        'document_ids': [first['id']], 'format': 'csv',
    }, headers={'Origin': 'https://testserver'}).status_code == 403


def test_selected_excel_has_readable_headers_and_a_usable_table(client, factory):
    from openpyxl import load_workbook

    first, second, third = cards(client, factory)
    response = client.post('/api/v1/documents/export-selection', json={
        'document_ids': [third['id'], first['id']], 'format': 'xlsx',
        'columns': ['title'], 'mode': 'editor',
    }, headers=headers(client))
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert list(sheet.values) == [('Наименование документа',), ('Selected third',), ('Selected first',)]
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref == 'A1:A3'
    assert sheet.column_dimensions['A'].width >= 45
    assert sheet['A1'].font.bold and sheet['A1'].fill.fgColor.rgb.endswith('147D68')
    assert sheet['A2'].alignment.wrap_text
    assert sheet['A1'].alignment.indent == 1 and sheet['A2'].alignment.indent == 1
    assert sheet['A2'].fill.fgColor.rgb != sheet['A3'].fill.fgColor.rgb
    assert sheet.sheet_view.showGridLines is False
    assert second['title'] not in str(list(sheet.values))


def test_filtered_excel_exports_all_matches_in_order_not_only_the_loaded_page(client, factory):
    import json

    from openpyxl import load_workbook

    from app.content.indexing import index_metadata

    documents = cards(client, factory)
    with factory.begin() as db:
        from uuid import UUID

        for document in documents:
            index_metadata(db, UUID(document['id']))
    parameters = {'query': 'Selected', 'scope': 'metadata', 'fuzzy': False,
                  'page_size': 1, 'sort': 'title', 'direction': 'desc', 'mode': 'editor'}
    result = client.post('/api/v1/search', json=parameters, headers=headers(client))
    assert result.status_code == 200, result.text
    assert result.json()['total'] == 2 and len(result.json()['items']) == 1
    response = client.get('/api/v1/search/export', params={
        'parameters': json.dumps(parameters), 'format': 'xlsx', 'columns': 'title',
    })
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert list(sheet.values) == [('Наименование документа',), ('Selected third',), ('Selected first',)]
    assert sheet.auto_filter.ref == 'A1:A3'
