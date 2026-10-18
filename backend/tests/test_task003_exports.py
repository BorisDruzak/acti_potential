import csv
import io

import pytest


def test_tabular_exports_preserve_numbers_and_neutralize_formulas():
    from openpyxl import load_workbook

    from app.registry.tabular import export_file

    values = ["0012", "=SUM(1,2)", " \t+CMD", "Русский текст"]
    content = export_file("csv", ["Номер", "Формула", "Пробел", "Текст"], [values])
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert rows[1] == ["0012", "'=SUM(1,2)", "' \t+CMD", "Русский текст"]
    book = load_workbook(io.BytesIO(export_file("xlsx", ["Номер"], [[v] for v in values])))
    assert book.active["A2"].value == "0012"
    assert book.active["A3"].value == "=SUM(1,2)"
    assert book.active["A3"].data_type == "s"


def test_tabular_import_rejects_duplicate_headers_and_oversized_rows():
    from app.registry.tabular import read_tabular

    with pytest.raises(ValueError):
        read_tabular("csv", b"code,code\na,b\n")
    with pytest.raises(ValueError):
        read_tabular("csv", b"code\n" + b"x\n" * 2001)
    assert read_tabular("csv", "code,label\n001,Тест\n".encode()) == [
        {"code": "001", "label": "Тест"}
    ]


def test_xlsx_import_accepts_native_dates_and_rejects_formula_cells():
    from datetime import datetime

    from openpyxl import Workbook

    from app.registry.tabular import read_tabular

    book = Workbook()
    book.active.append(["label", "valid_from"])
    book.active.append(["Native date", datetime(2026, 10, 15)])
    content = io.BytesIO()
    book.save(content)
    assert read_tabular("xlsx", content.getvalue()) == [
        {"label": "Native date", "valid_from": "2026-10-15"}
    ]
    book.active["A2"] = "=1+1"
    content = io.BytesIO()
    book.save(content)
    with pytest.raises(ValueError, match="Formulas"):
        read_tabular("xlsx", content.getvalue())


def test_document_export_and_ranges_enforce_row_and_field_acl(client, factory):
    import json

    from app.registry.models import Document, Registry
    from tests.test_content_security import actors_and_document, sign_in_as

    sample = actors_and_document(factory, fields=[
        {"key": "number", "type": "text", "system": "reg_number", "label": "Номер",
         "filterable": True, "exportable": True},
        {"key": "amount", "type": "decimal", "label": "Сумма",
         "filterable": True, "exportable": True},
        {"key": "secret", "type": "text", "label": "Скрыто", "read_roles": ["editor"],
         "filterable": True, "exportable": True},
        {"key": "no_export", "type": "text", "label": "Не экспортировать"},
    ], reg_number="0012")
    with factory.begin() as db:
        row = db.get(Document, sample["document"])
        db.get(Registry, sample["registry"]).document_type_ids = [str(row.document_type_id)]
        row.custom_data = {"amount": "12.50", "secret": "HIDDEN", "no_export": "NO-EXPORT"}
    sign_in_as(client, factory, sample["viewer"])
    params = {"filters": json.dumps({"amount": {"op": "between", "value": ["12", "13"]}})}
    result = client.get("/api/v1/documents", params=params)
    assert result.status_code == 200, result.text
    assert result.json()["total"] == 1
    assert result.json()["items"][0]["values"] == {
        "number": "0012", "amount": "12.50", "no_export": "NO-EXPORT"
    }
    exported = client.get("/api/v1/documents/export", params={**params, "format": "csv"})
    assert exported.status_code == 200, exported.text
    values = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))
    assert values == [["Номер [number]", "Сумма [amount]"], ["0012", "12.50"]]
    assert client.get("/api/v1/documents/export", params={"columns": "secret"}).status_code == 422
    assert client.get("/api/v1/documents", params={
        "filters": json.dumps({"secret": "HIDDEN"})
    }).status_code == 422
    with factory.begin() as db:
        from app.registry.models import DocumentACL

        db.add(DocumentACL(document_id=sample["document"], user_id=sample["editor"].id,
                           permissions=["read"]))
    assert len(list(csv.reader(io.StringIO(client.get(
        "/api/v1/documents/export", params={"format": "csv"}
    ).content.decode("utf-8-sig"))))) == 1


def test_export_checks_each_historical_schema_column_and_group_child(client, factory):
    from copy import deepcopy

    from openpyxl import load_workbook

    from app.registry.models import Document, FormSchema
    from tests.test_content_security import actors_and_document, sign_in_as

    fields = [
        {"key": "public", "type": "text", "label": "Public", "exportable": True},
        {"key": "rows", "type": "group", "label": "Rows", "exportable": True, "children": [
            {"key": "name", "type": "text", "label": "Name", "exportable": True},
            {"key": "secret", "type": "text", "label": "Secret", "exportable": True,
             "read_roles": ["editor"]},
            {"key": "excluded", "type": "text", "label": "Excluded"},
        ]},
    ]
    sample = actors_and_document(factory, fields=fields)
    with factory.begin() as db:
        first = db.get(Document, sample["document"])
        first.custom_data = {"public": "=SAFE-TEXT", "rows": [
            {"name": "Visible", "secret": "DO-NOT-EXPORT", "excluded": "EXCLUDED"}
        ]}
        schema = db.get(FormSchema, first.schema_version_id)
        newer = deepcopy(schema.schema_data)
        newer["sections"][0]["fields"][0]["exportable"] = False
        second_schema = FormSchema(document_type_id=first.document_type_id, version=2,
                                   state="published", schema_data=newer)
        db.add(second_schema)
        db.flush()
        db.add(Document(card_state="published", publication_targets=["internal", "shared"],
                        registry_id=first.registry_id, document_type_id=first.document_type_id,
                        schema_version_id=second_schema.id, created_by=sample["editor"].id,
                        custom_data={"public": "NEW-NONEXPORTABLE", "rows": []}))
    sign_in_as(client, factory, sample["viewer"])
    result = client.get("/api/v1/documents/export", params={"format": "xlsx", "columns": "public,rows"})
    assert result.status_code == 200, result.text
    rows = list(load_workbook(io.BytesIO(result.content)).active.values)
    serialized = str(rows)
    assert "DO-NOT-EXPORT" not in serialized and "EXCLUDED" not in serialized
    assert "NEW-NONEXPORTABLE" not in serialized
    assert "=SAFE-TEXT" in serialized and "Visible" in serialized


def test_numeric_and_date_presence_filters_do_not_cast_empty_text(client, factory):
    import json
    from datetime import date

    from app.registry.models import Document
    from tests.test_content_security import actors_and_document, sign_in_as

    sample = actors_and_document(factory, fields=[
        {"key": "date", "type": "date", "system": "reg_date", "filterable": True},
        {"key": "quantity", "type": "integer", "filterable": True},
    ], reg_date=date(2026, 1, 1))
    with factory.begin() as db:
        row = db.get(Document, sample["document"])
        row.custom_data = {"quantity": 5}
    sign_in_as(client, factory, sample["viewer"])
    for key in ("date", "quantity"):
        result = client.get("/api/v1/documents", params={"filters": json.dumps({key: {"op": "present"}})})
        assert result.status_code == 200 and result.json()["total"] == 1, result.text
    result = client.get("/api/v1/documents", params={"filters": json.dumps({
        "quantity": {"op": "between", "value": [1, "6"]}
    })})
    assert result.status_code == 200 and result.json()["total"] == 1, result.text


def test_malformed_filter_operators_and_deep_json_return_validation_errors(client, factory):
    import json

    from tests.test_content_security import actors_and_document, sign_in_as

    sample = actors_and_document(factory, fields=[
        {"key": "public", "type": "text", "filterable": True}
    ])
    sign_in_as(client, factory, sample["viewer"])
    for value in [json.dumps({"public": {"op": [], "value": "x"}}),
                  json.dumps({"public": {"op": {}, "value": "x"}}),
                  "[" * 1100 + "0" + "]" * 1100]:
        result = client.get("/api/v1/documents", params={"filters": value})
        assert result.status_code == 422
