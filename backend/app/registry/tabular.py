"""Bounded tabular interchange. Values never become executable spreadsheet cells."""

import csv
import io
import math
from datetime import date, datetime, time
from itertools import chain
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_CELL = 32767


def safe_text(value):
    value = "" if value is None else str(value)
    if len(value) > MAX_CELL or any(ord(c) < 32 and c not in "\t\r\n" for c in value):
        raise ValueError("Cell contains unsupported characters or exceeds 32767 characters")
    return value


def export_file(format, headers, rows, *, presentation=False, widths=None):
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        for row in chain([headers], rows):
            values = [safe_text(v) for v in row]
            writer.writerow([
                "'" + v if v.lstrip(" \t\r\n").startswith(("=", "+", "-", "@"))
                or v.startswith(("\t", "\r", "\n")) else v for v in values
            ])
        return output.getvalue().encode("utf-8-sig")
    if format != "xlsx":
        raise ValueError("Unsupported tabular format")
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell

    book = Workbook(write_only=True)
    sheet = book.create_sheet("Документы" if presentation else "ACTI")
    if presentation:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        widths = widths or [max(18, min(46, len(header) + 4)) for header in headers]
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_properties.tabColor = "147D68"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = "8"  # OOXML A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        body_font = Font(name="Calibri", size=11, color="203B46")
        header_fill = PatternFill("solid", fgColor="147D68")
        fills = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="F0F7F5")]
        alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
        header_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        border = Border(bottom=Side(style="hair", color="D9E7E3"))
    last_row = 1
    for row_index, row in enumerate(chain([headers], rows), 1):
        last_row = row_index
        values = list(row)
        if presentation:
            lines = max((sum(max(1, math.ceil(len(line) / max(1, widths[i] - 3)))
                             for line in safe_text(value).split("\n"))
                         for i, value in enumerate(values)), default=1)
            sheet.row_dimensions[row_index].height = min(409, max(36 if row_index == 1 else 30, lines * 15 + 12))
        cells = []
        for value in values:
            # Validate text even when writing native dates. Identifiers stay text;
            # formula-looking values are always inert string cells.
            text = safe_text(value)
            cell = WriteOnlyCell(sheet, value=value if presentation and isinstance(value, date) else text)
            if presentation and isinstance(value, date):
                cell.number_format = "DD.MM.YYYY HH:MM" if isinstance(value, datetime) else "DD.MM.YYYY"
            else:
                cell.data_type = "s"
                cell.number_format = "@"
            if presentation:
                cell.font = header_font if row_index == 1 else body_font
                cell.fill = header_fill if row_index == 1 else fills[row_index % 2]
                cell.alignment = header_alignment if row_index == 1 else alignment
                cell.border = border
            cells.append(cell)
        sheet.append(cells)
    if presentation:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"
        sheet.print_options.horizontalCentered = True
        sheet.oddFooter.center.text = "ACTI · Страница &P из &N"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def read_tabular(format, content):
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("Import exceeds 2 MiB")
    book = None
    try:
        if format == "csv":
            rows = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""))
        elif format == "xlsx":
            from openpyxl import load_workbook

            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 100 or sum(e.file_size for e in entries) > 16 * 1024 * 1024:
                    raise ValueError("Expanded workbook exceeds import limits")
            book = load_workbook(io.BytesIO(content), read_only=True, keep_links=False)
            sheet = book.active
            if (sheet.max_row or 0) > MAX_IMPORT_ROWS + 1 or (sheet.max_column or 0) > 30:
                raise ValueError("Workbook dimensions exceed import limits")

            def excel_rows():
                for row in sheet.iter_rows():
                    if any(c.data_type == "f" for c in row):
                        raise ValueError("Formulas cannot be imported")
                    yield [
                        c.value.date().isoformat()
                        if isinstance(c.value, datetime) and c.value.time() == time()
                        else c.value.isoformat() if isinstance(c.value, datetime) else c.value
                        for c in row
                    ]

            rows = excel_rows()
        else:
            raise ValueError("Unsupported tabular format")
        headers = [safe_text(v).strip() for v in next(rows, [])]
        if not headers or len(headers) > 30 or any(not v for v in headers):
            raise ValueError("Invalid headers")
        if len(headers) != len(set(headers)):
            raise ValueError("Duplicate headers")
        result = []
        for row in rows:
            if not any(v is not None and v != "" for v in row):
                continue
            if len(row) != len(headers) or len(result) >= MAX_IMPORT_ROWS:
                raise ValueError("Invalid row length or more than 2000 rows")
            result.append(dict(zip(headers, [safe_text(v) for v in row], strict=True)))
        return result
    except (BadZipFile, UnicodeError, csv.Error, KeyError, OSError, ParseError,
            DefusedXmlException) as error:
        raise ValueError("Invalid tabular file") from error
    finally:
        if book:
            book.close()
