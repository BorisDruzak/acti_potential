"""Foreground synthetic pilot. Configuration/receipts/reports belong outside Git.

Run from backend: python tools/task011_load_pilot.py --config PRIVATE_JSON preflight|seed|run|cleanup
Never point this tool at production. It accepts only its exact, fresh, UUID-named
test database and a separately provisioned empty Redis server. No service is installed.
Linux operator must launch the pilot in a transient aggregate RSS/CPU scope
(e.g. MemoryMax=1536M, CPUQuota=200%). RLIMIT_AS below limits virtual address
space per process, not aggregate RSS: native thread/allocator mappings need headroom.
"""

from __future__ import annotations

import argparse
import asyncio
import base64
import csv
import io
import ipaddress
import json
import math
import os
import secrets
import shutil
import signal
import socket
import ssl
import subprocess
import sys
import threading
import time
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path
from urllib.parse import urlsplit
from uuid import UUID, uuid5

GIB = 1024**3
MARKER = "task011:load:owner"
MAX_SECONDS = 1200
MAX_ADDRESS_SPACE = 3 * GIB
CURRENT_STAGE = "startup"


class PilotError(ValueError):
    """Only static, non-sensitive harness diagnostics may use this type."""


@dataclass(frozen=True)
class PilotConfig:
    run_id: str
    database_url: str = field(repr=False)
    expected_database_oid: int
    redis_url: str = field(repr=False)
    expected_redis_run_id: str = field(repr=False)
    origin: str = field(repr=False)
    tls_cert: str = field(repr=False)
    tls_key: str = field(repr=False)
    tls_ca: str = field(repr=False)
    storage_mount: str = field(repr=False)
    expected_storage_uuid: str = field(repr=False)
    db_mount: str = field(repr=False)
    expected_db_uuid: str = field(repr=False)
    cards: int = 1000
    clients: int = 32
    seconds: int = 10
    request_limit: int = 3000
    reindex_limit: int = 250
    sample_files: int = 4

    @property
    def identity(self):
        return UUID(self.run_id)

    @property
    def database_name(self):
        return "acti_task011_" + self.identity.hex + "_test"


def loopback(host):
    return host == "localhost" or bool(host and ipaddress.ip_address(host).is_loopback)


def validate_config(cfg):
    from sqlalchemy.engine import make_url

    database = make_url(cfg.database_url)
    if (
        database.drivername != "postgresql+psycopg"
        or not loopback(database.host)
        or database.database != cfg.database_name
        or database.query
        or database.port is None
        or database.port == 5432
    ):
        raise PilotError("Refusing database outside exact isolated pilot namespace")
    redis, origin = urlsplit(cfg.redis_url), urlsplit(cfg.origin)
    if (
        redis.scheme not in {"redis", "rediss"}
        or not loopback(redis.hostname)
        or redis.port is None
        or redis.port == 6379
        or redis.path != "/0"
        or redis.query
        or redis.fragment
    ):
        raise PilotError("Refusing Redis outside dedicated loopback server")
    if (
        origin.scheme != "https"
        or not loopback(origin.hostname)
        or origin.port is None
        or origin.port < 1024
        or origin.path not in {"", "/"}
        or origin.query
        or origin.fragment
        or origin.username
    ):
        raise PilotError("Verified loopback HTTPS with a dedicated port is required")
    for value, maximum in (
        (cfg.cards, 150000),
        (cfg.clients, 64),
        (cfg.seconds, 300),
        (cfg.request_limit, 100000),
        (cfg.reindex_limit, 1000),
    ):
        if type(value) is not int or not 1 <= value <= maximum:
            raise PilotError("Pilot resource budget exceeded")
    if (
        not 0 <= cfg.sample_files <= 16
        or cfg.expected_database_oid <= 0
        or not cfg.expected_redis_run_id
    ):
        raise PilotError("Pilot target fingerprints are required")
    if not cfg.expected_storage_uuid or not cfg.expected_db_uuid:
        raise PilotError("Expected data mount identities are required")
    if cfg.storage_mount != "/srv/acti/documents" or cfg.db_mount != "/srv/acti/postgresql":
        raise PilotError("Pilot must use the project's dedicated data volumes")


def guarded_storage(mount, run_id):
    mount = Path(mount).absolute()
    candidate = mount / ".task011-load" / UUID(str(run_id)).hex
    if mount.resolve() != mount or candidate.resolve() != candidate:
        raise PilotError("Symlinks are not allowed in pilot storage")
    if not candidate.is_relative_to(mount) or candidate == mount:
        raise PilotError("Invalid pilot storage boundary")
    return candidate


def assert_cleanup_owner(cfg, receipt, actual_oid, marker, users, registries):
    if (
        receipt.get("run_id") != cfg.run_id
        or receipt.get("database_oid") != actual_oid
        or actual_oid != cfg.expected_database_oid
        or marker != cfg.run_id
        or not set(users).issubset(receipt.get("user_ids", []))
        or not set(registries).issubset(receipt.get("registry_ids", []))
    ):
        raise PilotError("Cleanup ownership verification failed")


def assert_cleanup_resources(cfg, redis_run_id, redis_owner, redis_keys, storage_owner):
    if (
        redis_run_id != cfg.expected_redis_run_id
        or redis_owner != cfg.run_id
        or set(redis_keys) != {MARKER}
        or storage_owner != cfg.run_id
    ):
        raise PilotError("Cleanup resource ownership verification failed")


def write_private(path, data):
    path = Path(path)
    temporary = path.with_name(path.name + ".new")
    descriptor = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    temporary.replace(path)


def mount_guard(path, expected_uuid, minimum_free):
    resolved = Path(path).resolve()
    if resolved != Path(path).absolute() or not os.path.ismount(resolved):
        raise PilotError("Expected dedicated data mount is absent")
    result = subprocess.run(
        ["findmnt", "-n", "-o", "UUID", "--target", str(resolved)],
        capture_output=True,
        text=True,
        check=True,
        timeout=5,
    )
    if result.stdout.strip() != expected_uuid or shutil.disk_usage(resolved).free < minimum_free:
        raise PilotError("Mount identity or free-space budget verification failed")


def database_engine(cfg):
    from sqlalchemy import create_engine

    return create_engine(
        cfg.database_url,
        hide_parameters=True,
        pool_size=2,
        max_overflow=0,
        connect_args={
            "connect_timeout": 3,
            "options": "-c statement_timeout=15000 -c lock_timeout=2000",
        },
    )


def database_identity(connection, cfg):
    from sqlalchemy import text

    name, oid = connection.execute(
        text("SELECT current_database(), oid FROM pg_database WHERE datname=current_database()")
    ).one()
    if name != cfg.database_name or oid != cfg.expected_database_oid:
        raise PilotError("Database identity differs from provisioned pilot target")
    return oid


def schema_marker(connection):
    from sqlalchemy import text

    return connection.scalar(
        text("SELECT obj_description(oid,'pg_namespace') FROM pg_namespace WHERE nspname='public'")
    )


def assert_schema_claim(cfg, actual_oid, marker, tables, receipt_exists):
    if actual_oid != cfg.expected_database_oid:
        raise PilotError("Database ownership fingerprint differs")
    if marker == cfg.run_id:
        if not receipt_exists:
            raise PilotError("Owned database requires its original receipt")
        return
    if marker in {None, "standard public schema"} and not tables and not receipt_exists:
        return
    raise PilotError("Public schema is populated or belongs to another run")


def redis_connection(cfg, *, require_claim=False):
    from redis import Redis

    redis = Redis.from_url(
        cfg.redis_url, decode_responses=True, socket_timeout=3, socket_connect_timeout=3
    )
    if redis.info("server").get("run_id") != cfg.expected_redis_run_id:
        redis.close()
        raise PilotError("Redis fingerprint differs from provisioned server")
    owner = redis.get(MARKER)
    if require_claim and owner != cfg.run_id or owner not in {None, cfg.run_id}:
        redis.close()
        raise PilotError("Redis belongs to another run")
    keyspace = redis.info("keyspace")
    if any(name != "db0" and value.get("keys", 0) for name, value in keyspace.items()):
        redis.close()
        raise PilotError("Redis server is not isolated")
    if owner is None and redis.dbsize():
        redis.close()
        raise PilotError("Redis must be fresh before claiming")
    return redis


def preflight(cfg, receipt_path):
    validate_config(cfg)
    if sys.platform != "linux":
        raise PilotError("Actual pilot requires Linux process resource limits")
    mount_guard(cfg.storage_mount, cfg.expected_storage_uuid, GIB)
    mount_guard(cfg.db_mount, cfg.expected_db_uuid, 8 * GIB)
    storage = guarded_storage(cfg.storage_mount, cfg.identity)
    for name in (cfg.tls_cert, cfg.tls_key, cfg.tls_ca):
        if not Path(name).is_file():
            raise PilotError("Provisioned TLS material is absent")
    ssl.create_default_context(cafile=cfg.tls_ca)
    for record in socket.getaddrinfo(urlsplit(cfg.origin).hostname, None):
        if not ipaddress.ip_address(record[4][0]).is_loopback:
            raise PilotError("Loopback hostname resolves outside loopback")
    engine = database_engine(cfg)
    from sqlalchemy import inspect, text

    with engine.connect() as connection:
        oid = database_identity(connection, cfg)
        names = set(inspect(connection).get_table_names())
        marker = schema_marker(connection)
        assert_schema_claim(cfg, oid, marker, names, receipt_path.is_file())
        if marker == cfg.run_id:
            receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
            users = {str(value) for value in connection.scalars(text("SELECT id FROM users"))}
            registries = {
                str(value) for value in connection.scalars(text("SELECT id FROM registries"))
            }
            assert_cleanup_owner(cfg, receipt, cfg.expected_database_oid, marker, users, registries)
    engine.dispose()
    redis = redis_connection(cfg)
    redis.close()
    if storage.exists() and not (storage / ".owner").is_file():
        raise PilotError("Unclaimed storage directory already exists")
    if storage.exists() and (storage / ".owner").read_text(encoding="utf-8") != cfg.run_id:
        raise PilotError("Storage belongs to another run")
    return {"guards": "passed", "cards": cfg.cards, "clients": cfg.clients, "seconds": cfg.seconds}


def identity(cfg, kind):
    return uuid5(cfg.identity, kind)


def seed(cfg, receipt_path):
    preflight(cfg, receipt_path)
    engine = database_engine(cfg)
    from sqlalchemy import func, insert, select, text
    from sqlalchemy.orm import sessionmaker

    from app.content.models import DocumentIndexState, MetadataFragment
    from app.db.migrate import upgrade_database
    from app.identity.models import Role, User, UserRole
    from app.registry.models import (
        Document,
        DocumentACL,
        DocumentType,
        FormSchema,
        MetadataRevision,
        Registry,
        RegistryAccess,
    )
    from app.registry.validation import validate_schema

    storage = guarded_storage(cfg.storage_mount, cfg.identity)
    storage.mkdir(parents=True, exist_ok=True, mode=0o700)
    (storage / ".owner").write_text(cfg.run_id, encoding="utf-8")
    redis = redis_connection(cfg)
    redis.set(MARKER, cfg.run_id, nx=True)
    if redis.get(MARKER) != cfg.run_id:
        raise PilotError("Failed to claim Redis")
    redis.close()
    upgrade_database(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    receipt = (
        json.loads(receipt_path.read_text(encoding="utf-8")) if receipt_path.exists() else None
    )
    if receipt is None:
        with factory.begin() as session:
            if session.scalar(select(func.count()).select_from(User)) or session.scalar(
                select(func.count()).select_from(Document)
            ):
                raise PilotError("Fresh corpus guard failed")
            user_ids = [identity(cfg, f"user:{index}") for index in range(cfg.clients)]
            kind, form, registry = (
                identity(cfg, "type"),
                identity(cfg, "form"),
                identity(cfg, "registry"),
            )
            fields = [
                {
                    "key": "title",
                    "label": "Synthetic title",
                    "type": "text",
                    "system": "title",
                    "filterable": True,
                    "exportable": True,
                },
                {
                    "key": "cohort",
                    "label": "Synthetic cohort",
                    "type": "integer",
                    "filterable": True,
                    "exportable": True,
                },
                {"key": "body", "label": "Synthetic body", "type": "textarea", "exportable": True},
                {
                    "key": "secret_payload",
                    "label": "SYNTHETIC-HIDDEN-FIELD",
                    "type": "text",
                    "read_roles": ["editor", "admin"],
                },
                {"key": "original", "label": "Synthetic original", "type": "file"},
            ]
            session.add(DocumentType(id=kind, code="task011-pilot", name="Synthetic pilot type"))
            session.flush()
            session.add(
                FormSchema(
                    id=form,
                    document_type_id=kind,
                    version=1,
                    state="published",
                    schema_data=validate_schema(
                        {
                            "sections": [
                                {"key": "pilot", "title": "Synthetic pilot", "fields": fields}
                            ]
                        }
                    ),
                )
            )
            session.add(
                Registry(
                    id=registry,
                    code="task011-pilot",
                    name="Synthetic isolated load pilot",
                    document_type_ids=[str(kind)],
                )
            )
            session.add_all(
                User(id=user_id, display_name=f"Synthetic pilot actor {index}")
                for index, user_id in enumerate(user_ids)
            )
            session.flush()
            roles = {row.code: row.id for row in session.scalars(select(Role))}
            for index, user_id in enumerate(user_ids):
                session.add(
                    UserRole(
                        user_id=user_id, role_id=roles["editor" if index % 2 == 0 else "viewer"]
                    )
                )
                session.add(
                    RegistryAccess(
                        registry_id=registry,
                        user_id=user_id,
                        permissions=["read", "create", "update", "upload", "download"],
                    )
                )
            session.flush()
            registry_ids = [str(value) for value in session.scalars(select(Registry.id))]
            session.execute(
                text("SELECT set_config('application_name','acti-task011-synthetic',true)")
            )
            # UUID is parsed before interpolation; this is a marker, never a filesystem/DSN value.
            session.execute(text("COMMENT ON SCHEMA public IS '" + str(cfg.identity) + "'"))
        receipt = {
            "run_id": cfg.run_id,
            "database_oid": cfg.expected_database_oid,
            "user_ids": [str(value) for value in user_ids],
            "registry_ids": registry_ids,
            "registry_id": str(registry),
            "document_type_id": str(kind),
            "schema_id": str(form),
            "seeded_cards": 0,
            "sample_files_complete": False,
        }
        write_private(receipt_path, receipt)
    started = time.perf_counter()
    with factory() as session:
        existing = session.scalar(
            select(func.count())
            .select_from(Document)
            .where(Document.registry_id == UUID(receipt["registry_id"]))
        )
    if existing > cfg.cards:
        raise PilotError("Cannot shrink a corpus implicitly")
    for start in range(existing, cfg.cards, 1000):
        documents, revisions, acl, fragments, states = [], [], [], [], []
        for index in range(start, min(cfg.cards, start + 1000)):
            document_id = identity(cfg, f"document:{index}")
            title = f"Синтетический документ {index:06d}"
            values = {
                "title": title,
                "cohort": index // 2000,
                "body": "Синтетический регламент подготовки пилотного документа",
                "secret_payload": "SYNTHETIC-HIDDEN-OPERAND",
            }
            documents.append(
                {
                    "id": document_id,
                    "registry_id": UUID(receipt["registry_id"]),
                    "document_type_id": UUID(receipt["document_type_id"]),
                    "schema_version_id": UUID(receipt["schema_id"]),
                    "title": title,
                    "reg_date": date(2026, 1, 1),
                    "created_by": UUID(receipt["user_ids"][0]),
                    "custom_data": {key: value for key, value in values.items() if key != "title"},
                }
            )
            revisions.append(
                {
                    "id": identity(cfg, f"revision:{index}"),
                    "document_id": document_id,
                    "revision": 1,
                    "values": values,
                    "created_by": UUID(receipt["user_ids"][0]),
                }
            )
            states.append({"document_id": document_id, "document_revision": 1})
            for key in ("title", "body", "secret_payload"):
                fragments.append(
                    {
                        "id": identity(cfg, f"fragment:{index}:{key}"),
                        "document_id": document_id,
                        "document_revision": 1,
                        "field_key": key,
                        "label": "Synthetic " + key,
                        "text": values[key],
                    }
                )
            if index % 10 == 0:
                acl.append(
                    {
                        "document_id": document_id,
                        "user_id": UUID(receipt["user_ids"][0]),
                        "permissions": ["read", "update", "upload", "download"],
                    }
                )
        with factory.begin() as session:
            for model, rows in (
                (Document, documents),
                (MetadataRevision, revisions),
                (DocumentACL, acl),
                (MetadataFragment, fragments),
                (DocumentIndexState, states),
            ):
                if rows:
                    session.execute(insert(model), rows)
        receipt["seeded_cards"] = min(cfg.cards, start + 1000)
        write_private(receipt_path, receipt)
    with engine.begin() as connection:
        connection.execute(text("ANALYZE documents"))
        connection.execute(text("ANALYZE content_metadata_fragments"))
        database_bytes = connection.scalar(text("SELECT pg_database_size(current_database())"))
    engine.dispose()
    return {
        "seeded_cards": cfg.cards,
        "bulk_seed_seconds": round(time.perf_counter() - started, 3),
        "seed_batches": math.ceil(cfg.cards / 1000),
        "database_bytes": database_bytes,
    }


def percentile(values, fraction):
    return (
        round(sorted(values)[max(0, math.ceil(len(values) * fraction) - 1)], 2) if values else None
    )


def statistics_output(observations):
    return {
        name: {
            "requests": len(rows),
            "statuses": dict(Counter(str(status) for status, _ in rows)),
            "p50_ms": percentile([duration for _, duration in rows], 0.5),
            "p95_ms": percentile([duration for _, duration in rows], 0.95),
            "max_ms": round(max((duration for _, duration in rows), default=0), 2),
        }
        for name, rows in observations.items()
    }


def issue_pilot_sessions(factory, receipt):
    from app.identity.sessions import issue_session

    with factory.begin() as session:
        return [
            issue_session(session, UUID(value), datetime.now(UTC)) for value in receipt["user_ids"]
        ]


def server_environment(cfg):
    env = {key: value for key, value in os.environ.items() if not key.startswith("ACTI_")}
    env.update(
        ACTI_DATABASE_URL=cfg.database_url,
        ACTI_REDIS_URL=cfg.redis_url,
        ACTI_SESSION_SECRET=secrets.token_urlsafe(32),
        ACTI_ORIGIN=cfg.origin,
        ACTI_STORAGE_ROOT=str(guarded_storage(cfg.storage_mount, cfg.identity)),
        ACTI_SECURE_COOKIE="true",
        ACTI_SOURCE_SHA="task011-synthetic-pilot",
        PGOPTIONS="-c statement_timeout=15000 -c lock_timeout=2000",
        PYTHONUNBUFFERED="1",
    )
    return env


def launch_server(cfg):
    port = urlsplit(cfg.origin).port
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", port))  # Fail before Popen if another listener owns this port.
    command = [
        sys.executable,
        "-m",
        "uvicorn",
        "app.main:create_app",
        "--factory",
        "--host",
        "127.0.0.1",
        "--port",
        str(port),
        "--workers",
        "1",
        "--ssl-certfile",
        cfg.tls_cert,
        "--ssl-keyfile",
        cfg.tls_key,
        "--no-access-log",
        "--log-level",
        "error",
        "--limit-concurrency",
        "128",
        "--timeout-graceful-shutdown",
        "5",
    ]
    return subprocess.Popen(
        command,
        cwd=Path(__file__).resolve().parents[1],
        env=server_environment(cfg),
        stdin=subprocess.DEVNULL,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.PIPE,
    )


def copy_private_diagnostics(stream, descriptor, *, limit=1024 * 1024):
    written = 0
    try:
        while chunk := stream.read(4096):
            retained = chunk[: max(0, limit - written)]
            if retained:
                written += os.write(descriptor, retained)
    finally:
        os.close(descriptor)


def diagnostics_path(receipt_path):
    return receipt_path.with_name(
        receipt_path.name.removesuffix(".receipt.json") + ".run.stderr.log"
    )


def stop_server(process):
    if process.poll() is None:
        process.terminate()
        try:
            process.wait(timeout=8)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=3)


def reindex(cfg, receipt):
    from sqlalchemy import select
    from sqlalchemy.orm import sessionmaker

    from app.content.indexing import index_metadata
    from app.content.models import DocumentIndexState

    engine = database_engine(cfg)
    factory = sessionmaker(engine)
    started = time.perf_counter()
    limit = min(cfg.cards, cfg.reindex_limit)
    for start in range(0, limit, 10):
        with factory.begin() as session:
            for index in range(start, min(limit, start + 10)):
                document_id = identity(cfg, f"document:{index}")
                state = session.scalar(
                    select(DocumentIndexState)
                    .where(DocumentIndexState.document_id == document_id)
                    .with_for_update()
                )
                state.document_revision = 0
                session.flush()
                index_metadata(session, document_id)
    engine.dispose()
    return {
        "documents": limit,
        "batch_size": 10,
        "seconds": round(time.perf_counter() - started, 3),
    }


def docx_bytes(index):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/></Types>',
        )
        archive.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Синтетический оригинал пилота {index}</w:t></w:r></w:p></w:body></w:document>',
        )
    return stream.getvalue()


async def sample_originals(cfg, receipt, client, actor, factory):
    from app.content.worker import process_extractions
    from app.jobs.worker import process_batch

    if receipt.get("sample_files_complete"):
        return {"files": cfg.sample_files, "previously_seeded": True}
    started = time.perf_counter()
    indices = [index for index in range(cfg.cards) if index % 10][: cfg.sample_files]
    for index in indices:
        url = "/api/v1/documents/" + str(identity(cfg, f"document:{index}"))
        card = (await client.get(url)).json()
        changed = await client.put(
            url,
            headers={"Origin": cfg.origin, "X-CSRF-Token": actor.csrf_token},
            json={
                "revision": card["revision"],
                "values": {},
                "files": [
                    {
                        "field_key": "original",
                        "name": "synthetic.docx",
                        "content_base64": base64.b64encode(docx_bytes(index)).decode(),
                    }
                ],
            },
        )
        if changed.status_code != 200:
            raise PilotError("Synthetic original upload failed")
    await asyncio.to_thread(process_batch, factory, 100)
    for _ in indices:
        if (
            await asyncio.to_thread(
                process_extractions, factory, guarded_storage(cfg.storage_mount, cfg.identity)
            )
            != 1
        ):
            raise PilotError("Synthetic original extraction failed")
    receipt["sample_files_complete"] = True
    return {"files": len(indices), "seconds": round(time.perf_counter() - started, 3)}


async def run_http(cfg, receipt, receipt_path):
    global CURRENT_STAGE
    CURRENT_STAGE = "session_setup"
    import httpx
    from sqlalchemy.orm import sessionmaker

    engine = database_engine(cfg)
    factory = sessionmaker(engine, expire_on_commit=False)
    sessions = issue_pilot_sessions(factory, receipt)
    context = ssl.create_default_context(cafile=cfg.tls_ca)
    clients = [
        httpx.AsyncClient(
            base_url=cfg.origin,
            verify=context,
            trust_env=False,
            timeout=httpx.Timeout(30, connect=3),
            limits=httpx.Limits(max_connections=1, max_keepalive_connections=1),
            cookies={"acti_session": actor.token},
            follow_redirects=False,
        )
        for actor in sessions
    ]
    observations, failures = defaultdict(list), Counter()
    try:
        CURRENT_STAGE = "readiness"
        for _ in range(100):
            try:
                if (await clients[0].get("/api/v1/health/ready")).status_code == 200:
                    break
            except httpx.HTTPError:
                pass
            await asyncio.sleep(0.1)
        else:
            raise PilotError("Isolated HTTPS application did not become ready")
        CURRENT_STAGE = "originals"
        originals = await sample_originals(cfg, receipt, clients[0], sessions[0], factory)
        write_private(receipt_path, receipt)
        registry = receipt["registry_id"]
        cohort_size = min(2000, cfg.cards)
        viewer = clients[1] if len(clients) > 1 else clients[0]
        private_count = math.ceil(cohort_size / 10)
        export_result = {}
        CURRENT_STAGE = "exports"
        for format in ("csv", "xlsx"):
            start = time.perf_counter()
            response = await viewer.get(
                "/api/v1/documents/export",
                params={
                    "registry_id": registry,
                    "filters": json.dumps({"cohort": 0}),
                    "columns": "title,cohort",
                    "format": format,
                },
            )
            if response.status_code != 200 or b"SYNTHETIC-HIDDEN" in response.content:
                raise PilotError("Bounded export or its ACL check failed")
            if format == "csv":
                rows = len(list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))) - 1
            else:
                import openpyxl

                workbook = openpyxl.load_workbook(
                    io.BytesIO(response.content), read_only=True, data_only=True
                )
                rows = sum(1 for _ in workbook.active.iter_rows()) - 1
                workbook.close()
            expected = cohort_size - private_count if len(clients) > 1 else cohort_size
            if rows != expected:
                raise PilotError("Export disclosed an incorrect ACL-filtered count")
            export_result[format] = {
                "rows": rows,
                "bytes": len(response.content),
                "milliseconds": round((time.perf_counter() - start) * 1000, 2),
            }
        next_request = 0
        deadline = time.monotonic() + cfg.seconds
        private_id = str(identity(cfg, "document:0"))

        async def actor_load(actor_index):
            nonlocal next_request
            while time.monotonic() < deadline and next_request < cfg.request_limit:
                serial = next_request
                next_request += 1
                scenario = serial % 5
                params = {"registry_id": registry, "page_size": 25}
                if scenario == 0:
                    name, path = "list", "/api/v1/documents"
                elif scenario == 1:
                    name, path = "metadata_search", "/api/v1/documents"
                    params["q"] = "Синтетический"
                elif scenario == 2:
                    name, path = "cohort_filter", "/api/v1/documents"
                    params["filters"] = json.dumps({"cohort": 0})
                elif scenario == 3:
                    name, path, params = "private_document", "/api/v1/documents/" + private_id, {}
                else:
                    name, path = "content_search", "/api/v1/content/search"
                    params.update(
                        text="регламент", fuzzy="false", filters=json.dumps({"cohort": 0})
                    )
                start = time.perf_counter()
                try:
                    response = await clients[actor_index].get(path, params=params)
                    observations[name].append(
                        (response.status_code, (time.perf_counter() - start) * 1000)
                    )
                    expected_status = 404 if scenario == 3 and actor_index != 0 else 200
                    if response.status_code != expected_status:
                        failures[name + ":http"] += 1
                        continue
                    if scenario == 3:
                        continue
                    body = response.json()
                    size = cohort_size if scenario in {2, 4} else cfg.cards
                    expected = size if actor_index == 0 else size - math.ceil(size / 10)
                    if body["total"] != expected:
                        failures[name + ":acl_count"] += 1
                    if actor_index % 2 and "SYNTHETIC-HIDDEN" in response.text:
                        failures[name + ":field_acl"] += 1
                    if actor_index != 0 and private_id in response.text:
                        failures[name + ":document_acl"] += 1
                except (httpx.HTTPError, ValueError, KeyError):
                    observations[name].append((0, (time.perf_counter() - start) * 1000))
                    failures[name + ":transport_or_contract"] += 1

        started = time.perf_counter()
        CURRENT_STAGE = "load_and_reindex"
        workers = [asyncio.create_task(actor_load(index)) for index in range(len(clients))]
        reindex_task = asyncio.create_task(asyncio.to_thread(reindex, cfg, receipt))
        await asyncio.gather(*workers)
        reindexed = await reindex_task
        # Hidden indexed text must remain unsearchable to a reader, including after reindex.
        if len(clients) > 1:
            CURRENT_STAGE = "hidden_content_validation"
            hidden = await viewer.get(
                "/api/v1/content/search",
                params={
                    "registry_id": registry,
                    "text": "SYNTHETIC-HIDDEN-OPERAND",
                    "fuzzy": "false",
                    "filters": json.dumps({"cohort": 0}),
                },
            )
            if hidden.status_code != 200 or hidden.json()["total"] != 0:
                failures["hidden_content_acl"] += 1
        return {
            "cards": cfg.cards,
            "concurrent_sessions": len(clients),
            "load_seconds": round(time.perf_counter() - started, 3),
            "requests_started": next_request,
            "operations": statistics_output(observations),
            "export": export_result,
            "reindex": reindexed,
            "originals": originals,
            "failures": dict(failures),
            "status": "passed" if not failures else "failed",
            "scope": "synthetic isolated API sessions; metadata corpus bulk-seeded; no password/OIDC login or production OS privilege acceptance",
        }
    finally:
        await asyncio.gather(*(client.aclose() for client in clients))
        engine.dispose()


def run(cfg, receipt_path):
    preflight(cfg, receipt_path)
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    if receipt["seeded_cards"] != cfg.cards:
        raise PilotError("Seeded corpus does not match the requested scale")
    redis = redis_connection(cfg, require_claim=True)
    redis.close()
    private_log = diagnostics_path(receipt_path)
    descriptor = os.open(
        private_log,
        os.O_WRONLY | os.O_CREAT | os.O_TRUNC | os.O_NOFOLLOW | getattr(os, "O_BINARY", 0),
        0o600,
    )
    os.fchmod(descriptor, 0o600)
    process, pump, success = None, None, False
    try:
        process = launch_server(cfg)
        pump = threading.Thread(
            target=copy_private_diagnostics, args=(process.stderr, descriptor), daemon=True
        )
        pump.start()
        report = asyncio.run(run_http(cfg, receipt, receipt_path))
        success = report.get("status") == "passed"
        return report
    finally:
        if process is not None:
            stop_server(process)
        if pump is not None and pump.ident is not None:
            pump.join(timeout=3)
        else:
            os.close(descriptor)
        if process is not None:
            process.stderr.close()
        if success:
            private_log.unlink(missing_ok=True)


def cleanup(cfg, receipt_path):
    validate_config(cfg)
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    from sqlalchemy import inspect, select, text

    from app.db.base import Base
    from app.identity.models import User

    # Import the normal complete metadata before comparing unknown-table guards.
    from app.main import create_app  # noqa:F401
    from app.registry.models import Registry

    engine = database_engine(cfg)
    mount_guard(cfg.storage_mount, cfg.expected_storage_uuid, GIB)
    mount_guard(cfg.db_mount, cfg.expected_db_uuid, 8 * GIB)
    storage = guarded_storage(cfg.storage_mount, cfg.identity)
    if (storage / ".owner").read_text(encoding="utf-8") != cfg.run_id:
        raise PilotError("Storage receipt differs from this run")
    redis = redis_connection(cfg, require_claim=True)
    if redis.dbsize() != 1:
        redis.close()
        raise PilotError("Unknown Redis keys prevent cleanup")
    assert_cleanup_resources(
        cfg,
        redis.info("server")["run_id"],
        redis.get(MARKER),
        list(redis.scan_iter(count=2)),
        (storage / ".owner").read_text(encoding="utf-8"),
    )
    with engine.begin() as connection:
        oid = database_identity(connection, cfg)
        users = {str(row) for row in connection.scalars(select(User.id))}
        registries = {str(row) for row in connection.scalars(select(Registry.id))}
        assert_cleanup_owner(cfg, receipt, oid, schema_marker(connection), users, registries)
        if set(inspect(connection).get_table_names()) - (
            set(Base.metadata.tables) | {"alembic_version"}
        ):
            raise PilotError("Unknown tables prevent pilot cleanup")
        # This exact, receipt-owned disposable database was empty before seed.
        # Never issue this command against a generic *_test or the live database.
        connection.execute(text("DROP SCHEMA public CASCADE"))
        connection.execute(text("CREATE SCHEMA public"))
    engine.dispose()
    redis.delete(MARKER)
    redis.close()
    shutil.rmtree(storage)
    receipt_path.unlink()
    diagnostics_path(receipt_path).unlink(missing_ok=True)
    return {
        "cleanup": "passed",
        "boundary": "only exact receipt-owned disposable test schema, Redis marker and storage child",
    }


def resource_budget():
    import resource

    if os.getsid(0) != os.getpid():
        if os.getpgrp() == os.getpid():
            raise PilotError("Pilot process group belongs to another session")
        os.setsid()
    if os.getsid(0) != os.getpid() or os.getpgrp() != os.getpid():
        raise PilotError("An independently owned pilot process session is required")
    group = os.getpgrp()
    stopped = threading.Event()

    def hard_stop():
        if not stopped.wait(MAX_SECONDS):
            os.killpg(group, signal.SIGKILL)

    threading.Thread(target=hard_stop, daemon=True).start()
    resource.setrlimit(resource.RLIMIT_AS, (MAX_ADDRESS_SPACE, MAX_ADDRESS_SPACE))
    resource.setrlimit(resource.RLIMIT_NOFILE, (256, 256))
    available = sorted(os.sched_getaffinity(0))
    os.sched_setaffinity(0, available[:2])
    return stopped


def main():
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("mode", choices=("preflight", "seed", "run", "cleanup"))
    args = parser.parse_args()
    try:
        cfg = PilotConfig(**json.loads(args.config.read_text(encoding="utf-8")))
        validate_config(cfg)
        if sys.platform != "linux":
            raise PilotError("Actual pilot is supported only on Linux")
        import fcntl

        if args.config.stat().st_mode & 0o077:
            raise PilotError("Private pilot configuration must have mode 0600")
        lock = args.config.with_suffix(".lock").open("w")
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        stopped = resource_budget()
        receipt_path = args.config.with_suffix(".receipt.json")
        command = {"preflight": preflight, "seed": seed, "run": run, "cleanup": cleanup}[args.mode]
        report = command(cfg, receipt_path)
        report["resource_boundary"] = {
            "per_process_virtual_bytes": MAX_ADDRESS_SPACE,
            "cpu_affinity_max": 2,
            "whole_run_max_seconds": MAX_SECONDS,
            "aggregate_rss": "operator-provided transient scope; not enforced by RLIMIT_AS",
        }
        write_private(args.config.with_suffix("." + args.mode + ".report.json"), report)
        print(json.dumps(report, ensure_ascii=False))
        stopped.set()
        return 1 if report.get("status") == "failed" else 0
    except Exception as error:
        # No exception text, request bodies, credentials, filesystem paths or tokens.
        print(
            json.dumps(
                {
                    "status": "failed",
                    "phase": args.mode,
                    "stage": CURRENT_STAGE,
                    "error_type": type(error).__name__,
                    **({"error_code": str(error)} if isinstance(error, PilotError) else {}),
                }
            )
        )
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
