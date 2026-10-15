"""Generate synthetic Task005 fixtures without application or infrastructure access.

Validated generator dependencies (separate from production dependencies):
    python-docx==0.8.11 reportlab==4.5.1 pypdf==6.9.2

Example:
    python scripts/generate_task005_fixtures.py OUTPUT --soffice SOFFICE --font FONT.ttf

LibreOffice runs headless with its own temporary profile. On Windows pass its
console launcher soffice.com to --soffice; elsewhere pass the soffice executable.
No Windows COM is used. Existing
fixture files are preserved unless --overwrite is explicitly passed. Generated
binaries and LibreOffice profiles belong outside tracked source directories.
"""

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
import time
from importlib.metadata import version
from pathlib import Path
from uuid import uuid4

FIXTURES = (
    "public.docx",
    "private.docx",
    "public.pdf",
    "updated.pdf",
    "blank.pdf",
    "encrypted.pdf",
    "broken.pdf",
    "public.doc",
    "short.doc",
)


def convert_doc(soffice, source, destination, profile):
    destination.mkdir(parents=True)
    profile.mkdir(parents=True, exist_ok=True)
    command = [
        str(soffice),
        "-env:UserInstallation=" + profile.as_uri(),
        "--headless",
        "--convert-to",
        "doc:MS Word 97",
        "--outdir",
        str(destination),
        str(source),
    ]
    options = {}
    if sys.platform == "win32":
        startup = subprocess.STARTUPINFO()
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = subprocess.SW_HIDE
        options = {"startupinfo": startup, "creationflags": subprocess.CREATE_NO_WINDOW}
    completed = subprocess.run(
        command, shell=False, capture_output=True, timeout=60, **options
    )
    (destination / "conversion.log").write_bytes(completed.stdout + completed.stderr)
    if completed.returncode:
        raise RuntimeError("Headless LibreOffice conversion failed; inspect conversion.log")
    target = destination / (source.stem + ".doc")
    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        if target.is_file() and target.stat().st_size > 0:
            data = target.read_bytes()
            if data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                return target
        time.sleep(0.2)
    raise RuntimeError("Headless LibreOffice did not produce a Word97 OLE fixture")


def generate(output, soffice, font, *, overwrite=False):
    from docx import Document
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    output = output.expanduser().resolve()
    soffice, font = (
        soffice.expanduser().resolve(strict=True),
        font.expanduser().resolve(strict=True),
    )
    if not soffice.is_file() or not font.is_file():
        raise ValueError(
            "Explicit LibreOffice executable and TrueType font files are required"
        )
    output.mkdir(parents=True, exist_ok=True)
    if not overwrite and any(
        (output / name).exists() for name in (*FIXTURES, "fixtures.json")
    ):
        raise ValueError(
            "Output contains fixture files; use a fresh directory or --overwrite"
        )
    build = output / (".native-build-" + uuid4().hex)
    build.mkdir(mode=0o700)

    document = Document()
    document.add_heading("Синтетический цифровой архив", level=1)
    document.add_paragraph(
        "Синтетическая фраза о городском благоустройстве и цифровом архиве."
    )
    table = document.add_table(rows=2, cols=2)
    for row, values in enumerate(
        (("Показатель", "Значение"), ("Тестовая территория", "Зелёный квартал"))
    ):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    document.save(output / "public.docx")
    private = Document()
    private.add_paragraph("ЗАКРЫТЫЙФАЙЛМАРКЕР недоступное содержание скрытого поля.")
    private.save(output / "private.docx")
    long_document = Document(output / "public.docx")
    for index in range(60):
        long_document.add_paragraph(
            f"Синтетическая контрольная запись {index + 1}. Документ подготовлен только для "
            "проверки конвертера старого формата и не содержит реальных данных организации."
        )
    long_source = build / "old-doc-source.docx"
    long_document.save(long_source)

    pdfmetrics.registerFont(TTFont("SyntheticFixtureFont", str(font)))
    for name, text in (
        ("public.pdf", "Проверка цифрового PDF об организации дорожного движения"),
        ("updated.pdf", "Обновлённая редакция PDF КВАНТОВЫЙМАРШРУТ сохраняет версии"),
    ):
        pdf = canvas.Canvas(str(output / name), invariant=1)
        pdf.setFont("SyntheticFixtureFont", 12)
        pdf.drawString(40, 800, text)
        pdf.save()
    blank = canvas.Canvas(str(output / "blank.pdf"), invariant=1)
    blank.showPage()
    blank.save()
    encrypted = PdfWriter()
    encrypted.append(PdfReader(output / "public.pdf"))
    encrypted.encrypt("synthetic-password-only")
    encrypted.write(output / "encrypted.pdf")
    (output / "broken.pdf").write_bytes(b"%PDF-corrupted synthetic Task005 fixture")

    # A short OS-temporary path also avoids LibreOffice's Windows profile path limit.
    # Never reuse the interactive user's profile, even when the output path is short.
    with tempfile.TemporaryDirectory(prefix="acti005-lo-") as profile_directory:
        profile = Path(profile_directory).resolve()
        short_doc = convert_doc(soffice, output / "public.docx", build / "short", profile)
        long_doc = convert_doc(soffice, long_source, build / "long", profile)
    shutil.copyfile(short_doc, output / "short.doc")
    shutil.copyfile(long_doc, output / "public.doc")
    receipt = {
        "synthetic_only": True,
        "generator_dependencies": {
            name: version(name) for name in ("python-docx", "reportlab", "pypdf")
        },
        "files": [
            {
                "name": name,
                "bytes": (output / name).stat().st_size,
                "sha256": hashlib.sha256((output / name).read_bytes()).hexdigest(),
            }
            for name in FIXTURES
        ],
    }
    (output / "fixtures.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return receipt


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument(
        "output", type=Path, help="Directory for generated synthetic fixtures"
    )
    parser.add_argument(
        "--soffice", type=Path, required=True, help="Explicit LibreOffice executable"
    )
    parser.add_argument(
        "--font", type=Path, required=True, help="TrueType font with Cyrillic glyphs"
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Replace only the named generated fixture files",
    )
    args = parser.parse_args()
    receipt = generate(args.output, args.soffice, args.font, overwrite=args.overwrite)
    print(
        f"Generated {len(receipt['files'])} synthetic fixtures. No source systems were accessed."
    )


if __name__ == "__main__":
    main()
